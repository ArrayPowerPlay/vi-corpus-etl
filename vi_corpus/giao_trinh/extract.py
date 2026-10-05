"""
Xử lý thô giáo trình (Google Drive "Tổng hợp giáo trình Đại học"): kiểm kê file và trích text từng trang.

Đọc <data-root>/raw/giao_trinh/<ngành>/<môn>/<file> (bản rclone, không sửa), ghi
interim/giao_trinh_text/<sha256>.json cho mỗi file duy nhất (định dạng checkpoint và vòng lặp chạy lại
nằm ở vi_corpus.common.batch_extract, dùng chung với VJOL):

    {"sha256", "size", "paths": [đường dẫn, cùng nội dung], "nganh", "mon", "format",
     "pages": [{"text", "method"}], "error", "skipped", "created_at"}

Phần riêng của giáo trình ở file này (chiến lược ở docs/GIAO_TRINH.md, mục 2-3):
- extract_file: PDF qua vi_corpus.common.pdf_text.extract_pdf (trang scan thì OCR, mô hình nạp lười),
  pptx qua python-pptx (mỗi slide một trang), định dạng khác ghi skipped="unsupported_format".
- ngành = thư mục cấp 1, môn = thư mục cấp 2 dưới thư mục giáo trình.
"""

from collections.abc import Callable
from pathlib import Path

from vi_corpus.common.batch_extract import (  # noqa: F401 — inventory, file_sha256, make_ocr_getter: tên cũ vẫn import được từ đây
    file_sha256,
    inventory,
    make_ocr_getter,
    run_batch,
)
from vi_corpus.common.batch_extract import checkpoint_path as _checkpoint_path
from vi_corpus.common.batch_extract import status as _status
from vi_corpus.common.pdf_text import extract_pdf

TEXT_DIR = "interim/giao_trinh_text"
FORMATS = {".pdf": "pdf", ".pptx": "pptx"}


def checkpoint_path(data_root: Path, sha256: str) -> Path:
    """Đường dẫn checkpoint (file kết quả trích text) của một file giáo trình."""
    return _checkpoint_path(data_root, TEXT_DIR, sha256)


def extract_pptx(path: Path) -> list[dict]:
    """
    Trích text pptx: mỗi slide một trang gồm khung chữ (kể cả trong nhóm), bảng và ghi chú.

    Returns:
        [{"text", "method": "pptx"}] theo thứ tự slide.
    """
    from pptx import Presentation

    def shape_lines(shapes) -> list[str]:
        """Các dòng text của một danh sách shape (đệ quy vào nhóm shape)."""
        out = []
        for shape in shapes:
            if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
                out += shape_lines(shape.shapes)
            elif shape.has_text_frame:
                out += [p.text for p in shape.text_frame.paragraphs]
            elif shape.has_table:
                out += [" | ".join(c.text for c in row.cells) for row in shape.table.rows]
        return out

    pages = []
    for slide in Presentation(path).slides:
        lines = shape_lines(slide.shapes)
        if slide.has_notes_slide:
            lines += [p.text for p in slide.notes_slide.notes_text_frame.paragraphs]
        pages.append({"text": "\n".join(ln for ln in lines if ln.strip()), "method": "pptx"})
    return pages


def extract_file(path: Path, get_ocr: Callable[[], object | None]) -> dict:
    """
    Trích text một file theo định dạng (đuôi file); không ném lỗi ra ngoài.

    Returns:
        {"format", "pages", "error", "skipped"}: định dạng không hỗ trợ thì skipped="unsupported_format";
        file hỏng / có mật khẩu thì error là thông báo lỗi.
    """
    fmt = FORMATS.get(path.suffix.lower())
    out = {"format": fmt or path.suffix.lower().lstrip(".") or "unknown", "pages": [], "error": None, "skipped": None}
    if fmt is None:
        out["skipped"] = "unsupported_format"
        return out
    try:
        out["pages"] = extract_pdf(path, get_ocr) if fmt == "pdf" else extract_pptx(path)
    except Exception as exc:  # noqa: BLE001 — một file lỗi không được dừng cả lô
        out["error"] = f"{type(exc).__name__}: {exc}"[:300]
    return out


def run_extract(root: Path, data_root: Path, get_ocr: Callable[[], object | None],
                ocr_enabled: bool = True, limit_files: int | None = None) -> int:
    """
    Kiểm kê rồi trích text mọi file giáo trình chưa làm (xem vi_corpus.common.batch_extract.run_batch).

    Ngành = thư mục cấp 1, môn = thư mục cấp 2 dưới `root`.

    Returns:
        Số file bị lỗi trong lần chạy này.
    """
    def path_meta(path: Path) -> dict:
        """Ngành và môn từ đường dẫn tương đối so với `root`."""
        parts = path.relative_to(root).parts
        return {"nganh": parts[0] if len(parts) > 1 else "", "mon": parts[1] if len(parts) > 2 else ""}

    return run_batch(root, data_root, TEXT_DIR, extract_file, path_meta, get_ocr,
                     ocr_enabled=ocr_enabled, limit_files=limit_files, desc="Trích text giáo trình")


def status(root: Path, data_root: Path) -> dict:
    """Tóm tắt tiến độ giáo trình (xem vi_corpus.common.batch_extract.status)."""
    return _status(root, data_root, TEXT_DIR)
