"""
Trích text file Office: docx (python-docx), pptx (python-pptx), doc / ppt (LibreOffice headless đổi sang docx / pptx).

LibreOffice chạy với hồ sơ người dùng riêng trong thư mục tạm cho mỗi lần đổi, để nhiều tiến trình chạy song song không
tranh nhau khoá hồ sơ mặc định. Cài trên server: docs/README.md Phần E (Q5).
"""

import shutil
import subprocess
import tempfile
from pathlib import Path

from vi_corpus.common.parsers import MissingTool

CONVERT_TIMEOUT = 300  # giây cho một file


def extract_docx(path: Path) -> list[dict]:
    """
    Trích text docx: đoạn văn và bảng theo đúng thứ tự trong thân tài liệu, mỗi đoạn / hàng bảng một dòng.

    Returns:
        Một trang [{"text", "method": "docx"}] (docx không có khái niệm trang cố định).
    """
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    lines: list[str] = []
    for el in document.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            lines.append(Paragraph(el, document).text)
        elif tag == "tbl":
            for row in Table(el, document).rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
    return [{"text": "\n".join(ln for ln in lines if ln.strip()), "method": "docx"}]


def extract_pptx(path: Path) -> list[dict]:
    """
    Trích text pptx: mỗi slide một trang gồm khung chữ (kể cả trong nhóm), bảng và ghi chú.

    Returns:
        [{"text", "method": "pptx"}] theo thứ tự slide.
    """
    from pptx import Presentation

    def shape_lines(shapes) -> list[str]:
        """Các dòng text của một danh sách shape (đệ quy vào nhóm shape)."""
        out = []
        for shape in shapes:
            if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
                out += shape_lines(shape.shapes)
            elif shape.has_text_frame:
                out += [p.text for p in shape.text_frame.paragraphs]
            elif shape.has_table:
                out += [" | ".join(c.text for c in row.cells) for row in shape.table.rows]
        return out

    pages = []
    for slide in Presentation(path).slides:
        lines = shape_lines(slide.shapes)
        if slide.has_notes_slide:
            lines += [p.text for p in slide.notes_slide.notes_text_frame.paragraphs]
        pages.append({"text": "\n".join(ln for ln in lines if ln.strip()), "method": "pptx"})
    return pages


def office_binary() -> str | None:
    """Đường dẫn lệnh LibreOffice (soffice / libreoffice), None nếu chưa cài."""
    return shutil.which("soffice") or shutil.which("libreoffice")


def convert_with_libreoffice(path: Path, target: str, out_dir: Path) -> Path:
    """
    Đổi file doc / ppt sang `target` ("docx" / "pptx") bằng LibreOffice headless, ghi vào out_dir.

    Raises:
        MissingTool: chưa cài LibreOffice.
        RuntimeError: LibreOffice không tạo ra file (file hỏng, có mật khẩu, quá thời gian).
    """
    binary = office_binary()
    if binary is None:
        raise MissingTool("needs_libreoffice")
    profile = out_dir / "lo_profile"
    cmd = [binary, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--norestore", "--convert-to", target,
           "--outdir", str(out_dir), str(path)]
    try:
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=CONVERT_TIMEOUT, check=False)
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError(f"LibreOffice quá {CONVERT_TIMEOUT}s") from exc
    out = out_dir / f"{path.stem}.{target}"
    if not out.exists():
        raise RuntimeError(f"LibreOffice không đổi được file: {(proc.stderr or proc.stdout).strip()[:200]}")
    return out


def extract_legacy(path: Path, target: str) -> list[dict]:
    """
    Trích text doc / ppt: đổi sang docx / pptx bằng LibreOffice rồi đọc như file mới; method giữ tên định dạng mới.

    Raises:
        MissingTool: chưa cài LibreOffice.
    """
    with tempfile.TemporaryDirectory(prefix="vi_corpus_lo_") as tmp:
        converted = convert_with_libreoffice(path, target, Path(tmp))
        return extract_docx(converted) if target == "docx" else extract_pptx(converted)
